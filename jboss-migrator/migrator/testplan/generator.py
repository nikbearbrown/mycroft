"""Builds the per-class unit test plan from the step commits on the migration branch.

For each changed production class it records what changed, which step/recipe changed it, a
risk level, the existing tests that cover it, the test cases to run, and gaps that need a
hand-written test.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from .. import javasrc
from ..git_ops import Repo
from ..planner import Step
from ..verify import diff_guard

RISK_ORDER = ["low", "medium", "high"]
HIGH_RISK_ANNOTATIONS = {"Entity", "Table", "NamedQuery", "RolesAllowed", "SecurityDomain",
                         "DeclareRoles", "RunAs", "TransactionAttribute", "MessageDriven"}

STEP_TEST_CASES = {
    "java-upgrade": "Runs on Java {java}: no reflective-access or removed-API failures at runtime",
    "javax-to-jakarta": "Deploys on target: CDI/EJB injection and lifecycle callbacks still fire",
    "hibernate-5-to-6": "Persistence: entity mappings load and queries return the baseline results",
    "junit4-to-junit5": "Migrated test class runs under JUnit 5 with the same assertions",
    "custom-recipes": "Behaviour unchanged after organisation recipe rewrite",
}


@dataclass
class StepCommit:
    step: Step
    before: str
    after: str


@dataclass
class Change:
    step_id: str
    description: str
    recipes: list[str]
    added: int
    removed: int
    namespace_only: bool


@dataclass
class ClassPlan:
    path: str
    fqcn: str
    protected: bool
    changes: list[Change] = field(default_factory=list)
    risk: str = "low"
    existing_tests: list[str] = field(default_factory=list)
    test_cases: list[str] = field(default_factory=list)
    gaps: list[str] = field(default_factory=list)


@dataclass
class TestPlan:
    classes: list[ClassPlan]
    characterization_tests: list[str]

    @property
    def gap_count(self) -> int:
        return sum(len(c.gaps) for c in self.classes)


def _max_risk(*risks: str) -> str:
    return max(risks, key=RISK_ORDER.index)


def _count(diff: str) -> tuple[int, int]:
    added = sum(1 for l in diff.splitlines() if l.startswith("+") and not l.startswith("+++"))
    removed = sum(1 for l in diff.splitlines() if l.startswith("-") and not l.startswith("---"))
    return added, removed


def _test_index(root: Path) -> dict[str, str]:
    """{repo-relative test path: source} for every test file."""
    return {p.relative_to(root).as_posix(): p.read_text(encoding="utf-8", errors="replace")
            for p in javasrc.test_sources(root)}


def _existing_tests(cls: javasrc.JavaClass, tests: dict[str, str]) -> list[str]:
    names = {f"{cls.name}Test", f"{cls.name}Tests", f"{cls.name}IT", f"Test{cls.name}",
             f"{cls.name}TestCase"}
    ref = re.compile(rf"\b{re.escape(cls.name)}\b")
    hits = []
    for path, src in tests.items():
        stem = Path(path).stem
        if stem in names or (stem != f"{cls.name}CharacterizationTest" and
                             (cls.fqcn in src or (ref.search(src) and f"package {cls.package};" in src))):
            hits.append(path)
    return sorted(hits)


def _covered_methods(test_paths: list[str], tests: dict[str, str]) -> str:
    return "\n".join(tests[p] for p in test_paths).lower()


def build(repo: Repo, step_commits: list[StepCommit], protected: list[str], java_version: int,
          characterization: list[str]) -> TestPlan:
    root = repo.path
    tests = _test_index(root)
    plans: dict[str, ClassPlan] = {}

    for sc in step_commits:
        for path in repo.changed_files(sc.before, sc.after):
            if not path.endswith(".java") or "src/test/" in path:
                continue
            file = root / path
            if not file.exists():
                continue
            cls = javasrc.parse(file, path)
            if cls is None:
                continue
            plan = plans.setdefault(path, ClassPlan(path, cls.fqcn,
                                                    diff_guard.is_protected(path, protected)))
            diff = repo.file_diff(sc.before, sc.after, path)
            added, removed = _count(diff)
            ns_only = diff_guard.check_diff(path, diff) is None
            plan.changes.append(Change(sc.step.id, sc.step.description, sc.step.recipes,
                                       added, removed, ns_only))
            plan.risk = _max_risk(plan.risk, "low" if ns_only else sc.step.risk)

    for path, plan in plans.items():
        cls = javasrc.parse(root / path, path)
        if HIGH_RISK_ANNOTATIONS & cls.annotations and not all(c.namespace_only for c in plan.changes):
            plan.risk = "high"
        plan.existing_tests = _existing_tests(cls, tests)
        covered = _covered_methods(plan.existing_tests, tests)

        for change in plan.changes:
            case = STEP_TEST_CASES.get(change.step_id)
            if case:
                plan.test_cases.append(case.format(java=java_version))
        for method in cls.public_methods:
            plan.test_cases.append(f"{method}(): returns the same result as baseline for "
                                   f"representative inputs")
            if method.lower() not in covered:
                plan.gaps.append(f"{method}() has no existing test referencing it")
        if not plan.existing_tests:
            plan.gaps.insert(0, "No existing test class found")
        plan.test_cases = list(dict.fromkeys(plan.test_cases))

    ordered = sorted(plans.values(), key=lambda p: (-RISK_ORDER.index(p.risk), p.path))
    return TestPlan(ordered, characterization)


def write_xlsx(plan: TestPlan, dest: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Test plan"
    headers = ["Class", "File", "Risk", "Protected", "Changes (step: +/-)", "Recipes",
               "Existing tests", "Test case", "Gap / manual test needed"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    fills = {"high": "F8CBAD", "medium": "FFE699", "low": "C6E0B4"}
    for c in plan.classes:
        changes = "\n".join(f"{ch.step_id}: +{ch.added}/-{ch.removed}"
                            f"{' (namespace only)' if ch.namespace_only else ''}" for ch in c.changes)
        recipes = "\n".join(sorted({r for ch in c.changes for r in ch.recipes}))
        rows = max(len(c.test_cases), len(c.gaps), 1)
        for i in range(rows):
            ws.append([c.fqcn, c.path, c.risk, "yes" if c.protected else "no", changes, recipes,
                       "\n".join(c.existing_tests) or "none",
                       c.test_cases[i] if i < len(c.test_cases) else "",
                       c.gaps[i] if i < len(c.gaps) else ""])
            ws.cell(ws.max_row, 3).fill = PatternFill("solid", fgColor=fills[c.risk])
    widths = [40, 50, 8, 10, 30, 45, 45, 70, 50]
    for col, width in zip("ABCDEFGHI", widths):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"

    if plan.characterization_tests:
        cs = wb.create_sheet("Characterization tests")
        cs.append(["Generated test (fill inputs, record golden output on baseline)"])
        for t in plan.characterization_tests:
            cs.append([t])
        cs.column_dimensions["A"].width = 100

    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    return dest
