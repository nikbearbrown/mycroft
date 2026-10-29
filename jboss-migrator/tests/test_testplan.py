from openpyxl import load_workbook

from migrator import report
from migrator.config import load
from migrator.planner import Step
from migrator.testplan import characterization, generator
from tests.conftest import ROOT, write

SERVICE = """package com.company.orders.service;

import javax.ejb.Stateless;
import javax.inject.Inject;

@Stateless
public class OrderService {
    @Inject
    OrderRepository repo;

    public Order find(long id) { return repo.find(id); }

    public java.util.List<Order> search(String customer) { return repo.byCustomer(customer); }
}
"""

TEST = """package com.company.orders.service;

import org.junit.Test;

public class OrderServiceTest {
    @Test
    public void findReturnsOrder() { new OrderService(); }
}
"""


def test_plan_maps_changes_tests_and_gaps(git_repo, tmp_path):
    write(git_repo.path, "src/main/java/com/company/orders/service/OrderService.java", SERVICE)
    write(git_repo.path, "src/test/java/com/company/orders/service/OrderServiceTest.java", TEST)
    before = git_repo.commit_all("baseline")

    svc = git_repo.path / "src/main/java/com/company/orders/service/OrderService.java"
    svc.write_text(SERVICE.replace("javax.", "jakarta."), encoding="utf-8")
    after = git_repo.commit_all("jakarta")

    step = Step("javax-to-jakarta", "openrewrite", "javax -> jakarta",
                ["org.openrewrite.java.migrate.jakarta.JakartaEE10"])
    plan = generator.build(git_repo, [generator.StepCommit(step, before, after)], [], 17, [])

    assert len(plan.classes) == 1
    c = plan.classes[0]
    assert c.fqcn == "com.company.orders.service.OrderService"
    assert c.risk == "low" and c.changes[0].namespace_only
    assert c.existing_tests == ["src/test/java/com/company/orders/service/OrderServiceTest.java"]
    assert any("injection" in t for t in c.test_cases)
    assert c.gaps == ["search() has no existing test referencing it"]

    cfg = load(ROOT / "migration.yaml")
    md = report.render("test_plan.md.j2", cfg=cfg, plan=plan)
    assert "`com.company.orders.service.OrderService` — risk: **low**" in md

    xlsx = generator.write_xlsx(plan, tmp_path / "plan.xlsx")
    ws = load_workbook(xlsx)["Test plan"]
    assert ws["A2"].value == "com.company.orders.service.OrderService"


def test_characterization_skeletons_follow_junit_version(tmp_path):
    write(tmp_path, "src/main/java/com/company/orders/service/OrderService.java", SERVICE)
    written = characterization.generate(tmp_path, ["com.company.orders.service"])
    assert [p.relative_to(tmp_path).as_posix() for p in written] == [
        "src/test/java/com/company/orders/service/OrderServiceCharacterizationTest.java"]
    src = written[0].read_text()
    assert "import org.junit.Ignore;" in src            # no JUnit 5 in project -> JUnit 4
    assert "public void find_matchesBaseline()" in src
    assert "public void search_matchesBaseline()" in src
    assert characterization.generate(tmp_path, ["com.company.orders.service"]) == []
